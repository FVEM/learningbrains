import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_validation_workbook():
    wb = Workbook()
    
    # -------------------------------------------------------------
    # 1. SHEET 1: RESPUESTAS (Responses_Itinerario)
    # -------------------------------------------------------------
    ws_responses = wb.active
    ws_responses.title = "Respuestas_Itinerario"
    ws_responses.views.sheetView[0].showGridLines = True
    
    # Colors
    c_navy = "1E3A8A"      # Group 1: Metadata
    c_blue = "2563EB"      # Section 2
    c_indigo = "4F46E5"    # Section 3
    c_sky = "0284C7"       # Section 4
    c_purple = "7C3AED"    # Section 5
    c_teal = "0D9488"      # Section 6
    c_emerald = "059669"   # Section 7
    c_amber = "D97706"     # Calculated Averages
    c_slate = "475569"     # Qualitative

    font_group = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10)
    font_formula = Font(name="Calibri", size=10, bold=True)

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Group Header Definitions (Row 1)
    # (start_col, end_col, title, hex_color)
    group_defs = [
        (1, 11, "1. METADATOS Y PERFIL DEL EVALUADOR / METADATA & EVALUATOR PROFILE", c_navy),
        (12, 15, "2. OBJETO Y PERTINENCIA / PURPOSE & RELEVANCE (P1-P4)", c_blue),
        (16, 19, "3. ESTRUCTURA Y COHERENCIA / STRUCTURE & COHERENCE (S1-S4)", c_indigo),
        (20, 24, "4. UNIDADES DE APRENDIZAJE / LEARNING UNITS (U1-U5)", c_sky),
        (25, 28, "5. PRINCIPIOS ÉTICOS Y LEGALES / ETHICS & INCLUSION (E1-E4)", c_purple),
        (29, 32, "6. MARCO DE COMPETENCIAS / COMPETENCE FRAMEWORK (C1-C4)", c_teal),
        (33, 37, "7. USO PRÁCTICO Y TRANSFERIBILIDAD / PRACTICAL USE & TRANSFER (T1-T5)", c_emerald),
        (38, 44, "PROMEDIOS CALCULADOS (1-5) / CALCULATED AVERAGES", c_amber),
        (45, 47, "8. COMENTARIOS CUALITATIVOS / QUALITATIVE FEEDBACK", c_slate),
    ]

    for start_c, end_c, title, color in group_defs:
        fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        ws_responses.merge_cells(start_row=1, start_column=start_c, end_row=1, end_column=end_c)
        cell = ws_responses.cell(row=1, column=start_c, value=title)
        cell.font = font_group
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for c in range(start_c, end_c + 1):
            ws_responses.cell(row=1, column=c).fill = fill
            ws_responses.cell(row=1, column=c).border = thin_border
    
    ws_responses.row_dimensions[1].height = 28

    # Detailed Column Headers (Row 2)
    # (col_idx, code, description, group_color, width, align)
    col_defs = [
        # Metadata
        (1, "Timestamp", "Fecha y hora (UTC)", c_navy, 20, "center"),
        (2, "Campaign", "ID Campaña", c_navy, 18, "center"),
        (3, "Language", "Idioma", c_navy, 10, "center"),
        (4, "Name", "Nombre y Apellidos", c_navy, 24, "left"),
        (5, "Email", "Email", c_navy, 26, "left"),
        (6, "Organization", "Organización / Entidad", c_navy, 26, "left"),
        (7, "Country", "País", c_navy, 10, "center"),
        (8, "Profile", "Perfil Profesional", c_navy, 22, "left"),
        (9, "Other_Detail", "Detalle otro perfil", c_navy, 20, "left"),
        (10, "Experience", "Años Experiencia", c_navy, 16, "center"),
        (11, "AI_Level", "Nivel Experiencia IA", c_navy, 18, "center"),

        # Section 2: Purpose
        (12, "P1", "P1: Necesidades puesto trabajo", c_blue, 12, "center"),
        (13, "P2", "P2: Objetivos claros", c_blue, 12, "center"),
        (14, "P3", "P3: Destinatarios claros", c_blue, 12, "center"),
        (15, "P4", "P4: Relevancia RRHH/formadores", c_blue, 12, "center"),

        # Section 3: Structure
        (16, "S1", "S1: Estructura clara y navegable", c_indigo, 12, "center"),
        (17, "S2", "S2: División en 4 partes útil", c_indigo, 12, "center"),
        (18, "S3", "S3: Conexión marcos europeos", c_indigo, 12, "center"),
        (19, "S4", "S4: Progresión coherente", c_indigo, 12, "center"),

        # Section 4: Learning Units
        (20, "U1", "U1: Cobertura competencias", c_sky, 12, "center"),
        (21, "U2", "U2: Claridad enfoque unidades", c_sky, 12, "center"),
        (22, "U3", "U3: Resultados realistas", c_sky, 12, "center"),
        (23, "U4", "U4: Equilibrio concepto/práctica", c_sky, 12, "center"),
        (24, "U5", "U5: Adaptable a perfiles/empresas", c_sky, 12, "center"),

        # Section 5: Ethics & Inclusion
        (25, "E1", "E1: Integración ética/legal", c_purple, 12, "center"),
        (26, "E2", "E2: Adecuado a destinatarios", c_purple, 12, "center"),
        (27, "E3", "E3: Conexión selección herramientas", c_purple, 12, "center"),
        (28, "E4", "E4: Equidad e inclusión", c_purple, 12, "center"),

        # Section 6: Competence Framework
        (29, "C1", "C1: Estructura KSA adecuada", c_teal, 12, "center"),
        (30, "C2", "C2: Alineación con unidades", c_teal, 12, "center"),
        (31, "C3", "C3: Actitudes responsables IA", c_teal, 12, "center"),
        (32, "C4", "C4: Base para nuevos contenidos", c_teal, 12, "center"),

        # Section 7: Practical Use & Transfer
        (33, "T1", "T1: Guía para Programa Formativo", c_emerald, 12, "center"),
        (34, "T2", "T2: Guía para formadores útil", c_emerald, 12, "center"),
        (35, "T3", "T3: Hoja de ruta adaptación", c_emerald, 12, "center"),
        (36, "T4", "T4: Claridad agentes externos", c_emerald, 12, "center"),
        (37, "T5", "T5: Potencial contextos europeos", c_emerald, 12, "center"),

        # Calculated Averages
        (38, "Media_P", "Media Sec 2 (Objeto)", c_amber, 14, "center"),
        (39, "Media_S", "Media Sec 3 (Estructura)", c_amber, 14, "center"),
        (40, "Media_U", "Media Sec 4 (Unidades)", c_amber, 14, "center"),
        (41, "Media_E", "Media Sec 5 (Ética)", c_amber, 14, "center"),
        (42, "Media_C", "Media Sec 6 (Competencias)", c_amber, 14, "center"),
        (43, "Media_T", "Media Sec 7 (Transferencia)", c_amber, 14, "center"),
        (44, "Media_Global", "Media Global (1-5)", c_amber, 15, "center"),

        # Qualitative Feedback
        (45, "Q1_Strengths", "Q1: Puntos fuertes principales", c_slate, 40, "left"),
        (46, "Q2_Improvements", "Q2: Aspectos a revisar o mejorar", c_slate, 40, "left"),
        (47, "Q3_Missing_Topics", "Q3: Temas o competencias ausentes", c_slate, 40, "left"),
    ]

    ws_responses.row_dimensions[2].height = 36

    for col_idx, code, label, color, width, align in col_defs:
        fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        cell = ws_responses.cell(row=2, column=col_idx, value=f"{code}\n({label})" if code != label else code)
        cell.font = font_header
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws_responses.column_dimensions[col_letter].width = width

    # Freeze top 2 rows
    ws_responses.freeze_panes = "A3"

    # Insert a sample illustrative row to test formula syntax and styling
    sample_row = 3
    ws_responses.cell(row=sample_row, column=1, value="2026-09-22 14:30:00").alignment = Alignment(horizontal="center")
    ws_responses.cell(row=sample_row, column=2, value="itinerario-formativo").alignment = Alignment(horizontal="center")
    ws_responses.cell(row=sample_row, column=3, value="ES").alignment = Alignment(horizontal="center")
    ws_responses.cell(row=sample_row, column=4, value="Ejemplo Evaluador")
    ws_responses.cell(row=sample_row, column=5, value="evaluador@ejemplo.eu")
    ws_responses.cell(row=sample_row, column=6, value="Centro FP Industrial / Empresa")
    ws_responses.cell(row=sample_row, column=7, value="ES").alignment = Alignment(horizontal="center")
    ws_responses.cell(row=sample_row, column=8, value="vet_trainer")
    ws_responses.cell(row=sample_row, column=9, value="")
    ws_responses.cell(row=sample_row, column=10, value="11-20").alignment = Alignment(horizontal="center")
    ws_responses.cell(row=sample_row, column=11, value="moderate").alignment = Alignment(horizontal="center")

    # Sample Likert scores (Cols 12 to 37: 26 questions)
    sample_scores = [5, 4, 5, 4, 5, 4, 5, 4, 5, 5, 4, 4, 5, 5, 4, 5, 4, 5, 4, 4, 5, 4, 5, 4, 5, 5]
    for i, score in enumerate(sample_scores):
        c = ws_responses.cell(row=sample_row, column=12 + i, value=score)
        c.alignment = Alignment(horizontal="center")

    # Formulas for row 3:
    # Media P: cols L-O (12 to 15)
    ws_responses.cell(row=sample_row, column=38, value=f"=AVERAGE(L{sample_row}:O{sample_row})").number_format = '0.00'
    # Media S: cols P-S (16 to 19)
    ws_responses.cell(row=sample_row, column=39, value=f"=AVERAGE(P{sample_row}:S{sample_row})").number_format = '0.00'
    # Media U: cols T-X (20 to 24)
    ws_responses.cell(row=sample_row, column=40, value=f"=AVERAGE(T{sample_row}:X{sample_row})").number_format = '0.00'
    # Media E: cols Y-AB (25 to 28)
    ws_responses.cell(row=sample_row, column=41, value=f"=AVERAGE(Y{sample_row}:AB{sample_row})").number_format = '0.00'
    # Media C: cols AC-AF (29 to 32)
    ws_responses.cell(row=sample_row, column=42, value=f"=AVERAGE(AC{sample_row}:AF{sample_row})").number_format = '0.00'
    # Media T: cols AG-AK (33 to 37)
    ws_responses.cell(row=sample_row, column=43, value=f"=AVERAGE(AG{sample_row}:AK{sample_row})").number_format = '0.00'
    # Media Global: cols L-AK (12 to 37)
    ws_responses.cell(row=sample_row, column=44, value=f"=AVERAGE(L{sample_row}:AK{sample_row})").number_format = '0.00'

    for col_f in range(38, 45):
        cell_f = ws_responses.cell(row=sample_row, column=col_f)
        cell_f.alignment = Alignment(horizontal="center")
        cell_f.font = font_formula
        cell_f.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

    # Sample Qualitative
    ws_responses.cell(row=sample_row, column=45, value="Excelente estructuración de las 6 unidades formativas y gran aplicabilidad en PYMEs.").alignment = Alignment(vertical="top", wrap_text=True)
    ws_responses.cell(row=sample_row, column=46, value="Reforzar con ejemplos concretos de prompts en el módulo de herramientas industriales.").alignment = Alignment(vertical="top", wrap_text=True)
    ws_responses.cell(row=sample_row, column=47, value="Ninguno en particular, el itinerario es muy completo.").alignment = Alignment(vertical="top", wrap_text=True)

    for c in range(1, 48):
        cell = ws_responses.cell(row=sample_row, column=c)
        cell.border = thin_border
        if cell.font != font_formula:
            cell.font = font_data

    # Auto-filter over all columns
    ws_responses.auto_filter.ref = f"A2:AU{sample_row}"

    # -------------------------------------------------------------
    # 2. SHEET 2: CODEBOOK & PREGUNTAS (Codebook_Preguntas)
    # -------------------------------------------------------------
    ws_codebook = wb.create_sheet(title="Codebook_Preguntas")
    ws_codebook.views.sheetView[0].showGridLines = True

    # Title Banner
    ws_codebook.merge_cells("A1:E1")
    title_cell = ws_codebook["A1"]
    title_cell.value = "LEARNING BRAINS (ERASMUS+) - CODEBOOK Y DICCIONARIO DE VARIABLES"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_codebook.row_dimensions[1].height = 32

    # Headers for Questions Table
    headers_cb = ["Código", "Sección", "Pregunta en Inglés (Original)", "Pregunta en Español", "Escala / Tipo"]
    for c_idx, h_text in enumerate(headers_cb, 1):
        cell = ws_codebook.cell(row=3, column=c_idx, value=h_text)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        cell.alignment = Alignment(horizontal="center" if c_idx in [1, 5] else "left", vertical="center")
        cell.border = thin_border
    ws_codebook.row_dimensions[3].height = 24

    questions_data = [
        # Section 2
        ("P1", "Sección 2: Objeto y pertinencia", "The itinerary addresses relevant needs related to on-the-job training in industrial companies.", "El itinerario responde a necesidades relevantes relacionadas con el aprendizaje en el puesto de trabajo en empresas industriales.", "Likert (1-5)"),
        ("P2", "Sección 2: Objeto y pertinencia", "The overall aims of the itinerary are clear.", "Los objetivos generales del itinerario son claros.", "Likert (1-5)"),
        ("P3", "Sección 2: Objeto y pertinencia", "The intended target groups are clearly identified.", "Los grupos destinatarios previstos están claramente identificados.", "Likert (1-5)"),
        ("P4", "Sección 2: Objeto y pertinencia", "The itinerary is relevant to HR and training managers and identified target groups.", "El itinerario es relevante para los responsables de RRHH, formadores y los grupos destinatarios identificados.", "Likert (1-5)"),

        # Section 3
        ("S1", "Sección 3: Estructura y coherencia", "The document is clearly structured and easy to navigate.", "El documento está claramente estructurado y es fácil de navegar y consultar.", "Likert (1-5)"),
        ("S2", "Sección 3: Estructura y coherencia", "The division of the document into four parts is useful and logical.", "La división del documento en cuatro partes es útil y lógica.", "Likert (1-5)"),
        ("S3", "Sección 3: Estructura y coherencia", "The relationship between the framework, the learning units, the guidance for trainers and the European reference frameworks is clear.", "La relación entre el marco, las unidades de aprendizaje, la guía para formadores y los marcos de referencia europeos es clara.", "Likert (1-5)"),
        ("S4", "Sección 3: Estructura y coherencia", "The progression between the learning units is coherent.", "La progresión entre las unidades de aprendizaje es coherente.", "Likert (1-5)"),

        # Section 4
        ("U1", "Sección 4: Unidades de aprendizaje", "The six learning units collectively cover the main competences needed to use AI in on-the-job training.", "Las seis unidades de aprendizaje cubren en conjunto las principales competencias necesarias para utilizar la IA en la formación en el puesto de trabajo.", "Likert (1-5)"),
        ("U2", "Sección 4: Unidades de aprendizaje", "The descriptions of the core focus of each unit are clear.", "Las descripciones del enfoque central de cada unidad formativa son claras.", "Likert (1-5)"),
        ("U3", "Sección 4: Unidades de aprendizaje", "The main learning outcomes are realistic and appropriate.", "Los principales resultados de aprendizaje previstos son realistas y adecuados.", "Likert (1-5)"),
        ("U4", "Sección 4: Unidades de aprendizaje", "The balance between conceptual understanding and practical application is appropriate.", "El equilibrio entre la comprensión conceptual y la aplicación práctica es el adecuado.", "Likert (1-5)"),
        ("U5", "Sección 4: Unidades de aprendizaje", "The learning units are sufficiently adaptable to different learner profiles and company contexts.", "Las unidades de aprendizaje son suficientemente adaptables a distintos perfiles de alumnado y contextos de empresa.", "Likert (1-5)"),

        # Section 5
        ("E1", "Sección 5: Principios éticos y legales", "Ethical and legal aspects are sufficiently integrated into the itinerary.", "Los aspectos éticos y legales están suficientemente integrados en el itinerario.", "Likert (1-5)"),
        ("E2", "Sección 5: Principios éticos y legales", "The treatment of ethical and legal issues is appropriate for the target groups.", "El tratamiento de las cuestiones éticas y legales es adecuado para los grupos destinatarios.", "Likert (1-5)"),
        ("E3", "Sección 5: Principios éticos y legales", "The ethical dimension is connected meaningfully with tool selection, training design and implementation.", "La dimensión ética se conecta de forma significativa con la selección de herramientas, el diseño formativo y su ejecución.", "Likert (1-5)"),
        ("E4", "Sección 5: Principios éticos y legales", "The itinerary gives appropriate attention to fairness, accessibility and inclusion.", "El itinerario presta la debida atención a la equidad, la accesibilidad y la inclusión.", "Likert (1-5)"),

        # Section 6
        ("C1", "Sección 6: Marco de competencias", "The Knowledge, Skills and Attitudes structure is appropriate for the itinerary.", "La estructura de Conocimientos, Habilidades y Actitudes (KSA) es adecuada para el itinerario.", "Likert (1-5)"),
        ("C2", "Sección 6: Marco de competencias", "The competence framework is aligned with the learning units.", "El marco de competencias está debidamente alineado con las unidades de aprendizaje.", "Likert (1-5)"),
        ("C3", "Sección 6: Marco de competencias", "The attitudes identified are appropriate for responsible and practical use of AI.", "Las actitudes identificadas son adecuadas para un uso responsable y práctico de la IA.", "Likert (1-5)"),
        ("C4", "Sección 6: Marco de competencias", "The competence framework can support the subsequent development of learning content and activities.", "El marco de competencias puede sustentar con éxito el posterior desarrollo de contenidos y actividades formativas.", "Likert (1-5)"),

        # Section 7
        ("T1", "Sección 7: Uso práctico y transferencia", "The itinerary provides useful guidance for developing a complete Training Programme.", "El itinerario proporciona directrices útiles para el desarrollo de un Programa Formativo integral.", "Likert (1-5)"),
        ("T2", "Sección 7: Uso práctico y transferencia", "The practical guidance for trainers is useful.", "Las orientaciones prácticas para formadores son de gran utilidad.", "Likert (1-5)"),
        ("T3", "Sección 7: Uso práctico y transferencia", "The adaptation roadmap can help external trainers adapt the itinerary to different learner groups.", "La hoja de ruta de adaptación puede ayudar a formadores externos a adaptar el itinerario a diversos grupos de participantes.", "Likert (1-5)"),
        ("T4", "Sección 7: Uso práctico y transferencia", "The itinerary is sufficiently clear to be used by trainers or organisations outside the consortium or beyond the project activities.", "El itinerario es suficientemente claro como para ser utilizado por formadores u organizaciones ajenas al consorcio o más allá de las actividades del proyecto.", "Likert (1-5)"),
        ("T5", "Sección 7: Uso práctico y transferencia", "The document has potential to support the development of training actions in different European contexts.", "El documento tiene un alto potencial para respaldar el desarrollo de acciones formativas en diversos contextos europeos.", "Likert (1-5)"),

        # Qualitative
        ("Q1_Strengths", "Sección 8: Cualitativa", "What do you consider to be the main strengths of the Learning Brains Training Itinerary?", "¿Cuáles consideras que son los principales puntos fuertes del Itinerario Formativo de Learning Brains?", "Texto Abierto"),
        ("Q2_Improvements", "Sección 8: Cualitativa", "What aspects, if any, would you suggest revising or developing further?", "¿Qué aspectos, en su caso, sugerirías revisar o desarrollar más a fondo?", "Texto Abierto"),
        ("Q3_Missing_Topics", "Sección 8: Cualitativa", "Do you think any relevant topic is missing from the Learning Brains Training Itinerary? If yes, please specify.", "¿Consideras que falta algún tema relevante en el Itinerario Formativo de Learning Brains? En caso afirmativo, indícalo por favor.", "Texto Abierto"),
    ]

    for idx, (code, sec, en_t, es_t, scale) in enumerate(questions_data, 4):
        ws_codebook.cell(row=idx, column=1, value=code).alignment = Alignment(horizontal="center")
        ws_codebook.cell(row=idx, column=2, value=sec)
        ws_codebook.cell(row=idx, column=3, value=en_t)
        ws_codebook.cell(row=idx, column=4, value=es_t)
        ws_codebook.cell(row=idx, column=5, value=scale).alignment = Alignment(horizontal="center")
        for c in range(1, 6):
            cell = ws_codebook.cell(row=idx, column=c)
            cell.border = thin_border
            cell.font = font_data

    ws_codebook.column_dimensions["A"].width = 15
    ws_codebook.column_dimensions["B"].width = 30
    ws_codebook.column_dimensions["C"].width = 45
    ws_codebook.column_dimensions["D"].width = 45
    ws_codebook.column_dimensions["E"].width = 18

    # -------------------------------------------------------------
    # 3. SHEET 3: KPI DASHBOARD (KPI_Dashboard)
    # -------------------------------------------------------------
    ws_kpi = wb.create_sheet(title="KPI_Dashboard")
    ws_kpi.views.sheetView[0].showGridLines = True

    ws_kpi.merge_cells("A1:D1")
    kpi_title = ws_kpi["A1"]
    kpi_title.value = "LEARNING BRAINS - RESUMEN EJECUTIVO DE VALIDACIÓN (INDICADORES)"
    kpi_title.font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    kpi_title.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    kpi_title.alignment = Alignment(horizontal="center", vertical="center")
    ws_kpi.row_dimensions[1].height = 30

    # Summary table 1: Medias por Sección
    sec_headers = ["Dimensión Evaluada", "Métrica", "Rango", "Fórmula"]
    for c_idx, h in enumerate(sec_headers, 1):
        c = ws_kpi.cell(row=3, column=c_idx, value=h)
        c.font = font_header
        c.fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        c.alignment = Alignment(horizontal="center" if c_idx > 1 else "left")
        c.border = thin_border

    dim_rows = [
        ("Total de Evaluaciones Recibidas", "=COUNTA(Respuestas_Itinerario!A3:A500)", "Nº respuestas", "Recuento"),
        ("Media Sección 2: Objeto y Pertinencia", "=AVERAGE(Respuestas_Itinerario!AL3:AL500)", "1 - 5", "Promedio P1-P4"),
        ("Media Sección 3: Estructura y Coherencia", "=AVERAGE(Respuestas_Itinerario!AM3:AM500)", "1 - 5", "Promedio S1-S4"),
        ("Media Sección 4: Unidades de Aprendizaje", "=AVERAGE(Respuestas_Itinerario!AN3:AN500)", "1 - 5", "Promedio U1-U5"),
        ("Media Sección 5: Principios Éticos y Legales", "=AVERAGE(Respuestas_Itinerario!AO3:AO500)", "1 - 5", "Promedio E1-E4"),
        ("Media Sección 6: Marco de Competencias", "=AVERAGE(Respuestas_Itinerario!AP3:AP500)", "1 - 5", "Promedio C1-C4"),
        ("Media Sección 7: Uso Práctico y Transferencia", "=AVERAGE(Respuestas_Itinerario!AQ3:AQ500)", "1 - 5", "Promedio T1-T5"),
        ("MEDIA GLOBAL DEL ITINERARIO FORMATIVO", "=AVERAGE(Respuestas_Itinerario!AR3:AR500)", "1 - 5", "Promedio Global (26 ítems)"),
    ]

    for idx, (dim, formula, rng, desc) in enumerate(dim_rows, 4):
        ws_kpi.cell(row=idx, column=1, value=dim).font = font_data
        val_cell = ws_kpi.cell(row=idx, column=2, value=formula)
        val_cell.font = font_formula
        val_cell.alignment = Alignment(horizontal="center")
        val_cell.number_format = '0.00' if idx > 4 else '0'
        ws_kpi.cell(row=idx, column=3, value=rng).alignment = Alignment(horizontal="center")
        ws_kpi.cell(row=idx, column=4, value=desc).alignment = Alignment(horizontal="left")

        if idx == 11: # Highlight overall mean
            for c in range(1, 5):
                ws_kpi.cell(row=idx, column=c).fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
                ws_kpi.cell(row=idx, column=c).font = Font(name="Calibri", size=10, bold=True, color="92400E")

        for c in range(1, 5):
            ws_kpi.cell(row=idx, column=c).border = thin_border

    ws_kpi.column_dimensions["A"].width = 42
    ws_kpi.column_dimensions["B"].width = 16
    ws_kpi.column_dimensions["C"].width = 16
    ws_kpi.column_dimensions["D"].width = 30

    # Save to file
    out_dir = os.path.join("public", "documents", "validation")
    os.makedirs(out_dir, exist_ok=True)
    out_file = os.path.join(out_dir, "Learning_Brains_Training_Itinerary_Validation_Responses.xlsx")
    wb.save(out_file)
    print(f"Generated: {out_file}")

if __name__ == "__main__":
    create_validation_workbook()
