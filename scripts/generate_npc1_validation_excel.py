import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_npc1_workbook():
    wb = Workbook()
    
    # -------------------------------------------------------------
    # 1. SHEET 1: VALIDATION RESPONSES (Validation_Responses)
    # -------------------------------------------------------------
    ws_responses = wb.active
    ws_responses.title = "Validation_Responses"
    ws_responses.views.sheetView[0].showGridLines = True
    
    # Professional Palette
    c_navy = "1E3A8A"      # Group 1: Metadata
    c_blue = "2563EB"      # Section 2: R1 Purpose & Relevance (P1-P6)
    c_indigo = "4F46E5"    # Section 3: Quality & Coverage (Q1-Q7)
    c_teal = "0D9488"      # Section 4: AI Tools, Trends & Cases (T1-T6)
    c_emerald = "059669"   # Section 5: Practical Value & Transferability (V1-V6)
    c_amber = "D97706"     # Calculated Averages
    c_slate = "475569"     # Qualitative Recommendations

    font_group = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10)
    font_formula = Font(name="Calibri", size=10, bold=True)

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Group Header Definitions (Row 1) - ALL IN ENGLISH
    group_defs = [
        (1, 10, "1. METADATA & EVALUATOR PROFILE", c_navy),
        (11, 16, "2. R1 PURPOSE & RELEVANCE (P1-P6)", c_blue),
        (17, 23, "3. QUALITY & COVERAGE OF MAPPING (Q1-Q7)", c_indigo),
        (24, 29, "4. AI TOOLS, TRENDS & CASE STUDIES (T1-T6)", c_teal),
        (30, 35, "5. PRACTICAL VALUE & TRANSFERABILITY (V1-V6)", c_emerald),
        (36, 40, "CALCULATED SECTION AVERAGES (1-5)", c_amber),
        (41, 45, "6. KEY RECOMMENDATIONS (QUALITATIVE)", c_slate),
    ]

    for start_c, end_c, title, color in group_defs:
        fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        ws_responses.merge_cells(start_row=1, start_column=start_c, end_row=1, end_column=end_c)
        cell = ws_responses.cell(row=1, column=start_c, value=title)
        cell.font = font_group
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for c in range(start_c, end_c + 1):
            ws_responses.cell(row=1, column=c).fill = fill
            ws_responses.cell(row=1, column=c).border = thin_border
    
    ws_responses.row_dimensions[1].height = 28

    # Detailed Column Headers (Row 2) - Exactly 45 Columns
    col_defs = [
        # Metadata (Cols 1-10)
        (1, "Timestamp", "Submission Date (UTC)", c_navy, 20, "center"),
        (2, "Campaign", "Campaign ID (npc-1)", c_navy, 16, "center"),
        (3, "Language", "Form Language", c_navy, 12, "center"),
        (4, "Full_Name", "Evaluator Name", c_navy, 24, "left"),
        (5, "Email", "Evaluator Email", c_navy, 26, "left"),
        (6, "Organization", "Organisation / Institution", c_navy, 28, "left"),
        (7, "Position", "Position / Role", c_navy, 24, "left"),
        (8, "Country", "Country", c_navy, 14, "center"),
        (9, "Professional_Profile", "Professional Profile", c_navy, 25, "left"),
        (10, "Other_Profile_Detail", "Other Profile Details", c_navy, 22, "left"),

        # Section 2: R1 Purpose and relevance (P1-P6) (Cols 11-16)
        (11, "P1", "P1: Purpose and intended use clear", c_blue, 13, "center"),
        (12, "P2", "P2: Addresses industrial training needs", c_blue, 13, "center"),
        (13, "P3", "P3: Relevant to HR, trainers & target groups", c_blue, 13, "center"),
        (14, "P4", "P4: Useful overview of AI in workplace learning", c_blue, 13, "center"),
        (15, "P5", "P5: Appropriate focus on industrial SMEs", c_blue, 13, "center"),
        (16, "P6", "P6: Relevant to support Training Programme", c_blue, 13, "center"),

        # Section 3: Quality and coverage of the mapping (Q1-Q7) (Cols 17-23)
        (17, "Q1", "Q1: Structured & easy to navigate", c_indigo, 13, "center"),
        (18, "Q2", "Q2: Findings & trends clear & understandable", c_indigo, 13, "center"),
        (19, "Q3", "Q3: Range of AI tools appropriate", c_indigo, 13, "center"),
        (20, "Q4", "Q4: Balanced opportunities, limits & challenges", c_indigo, 13, "center"),
        (21, "Q5", "Q5: Information current & relevant", c_indigo, 13, "center"),
        (22, "Q6", "Q6: Useful range of contexts (company, sectoral, national)", c_indigo, 13, "center"),
        (23, "Q7", "Q7: Conclusions consistent with evidence", c_indigo, 13, "center"),

        # Section 4: AI tools, trends and case studies (T1-T6) (Cols 24-29)
        (24, "T1", "T1: Explains strengths & limitations of AI apps", c_teal, 13, "center"),
        (25, "T2", "T2: Case studies clear & useful", c_teal, 13, "center"),
        (26, "T3", "T3: Illustrates realistic applications in workplace", c_teal, 13, "center"),
        (27, "T4", "T4: Information on suitability for org contexts", c_teal, 13, "center"),
        (28, "T5", "T5: Trends relevant to inform training content", c_teal, 13, "center"),
        (29, "T6", "T6: Distinguishes promising uses from limited ones", c_teal, 13, "center"),

        # Section 5: Practical value and transferability (V1-V6) (Cols 30-35)
        (30, "V1", "V1: Helps identify AI uses in own organisation", c_emerald, 13, "center"),
        (31, "V2", "V2: Transferable to industrial company contexts", c_emerald, 13, "center"),
        (32, "V3", "V3: Reference points for different digital maturity", c_emerald, 13, "center"),
        (33, "V4", "V4: Supports informed decision-making", c_emerald, 13, "center"),
        (34, "V5", "V5: Ethical, legal & inclusion considerations", c_emerald, 13, "center"),
        (35, "V6", "V6: Solid knowledge base for Programme & Toolkit", c_emerald, 13, "center"),

        # Calculated Averages (Cols 36-40)
        (36, "Mean_Sec2", "Mean Sec 2: Purpose & Relevance (P1-P6)", c_amber, 16, "center"),
        (37, "Mean_Sec3", "Mean Sec 3: Quality & Coverage (Q1-Q7)", c_amber, 16, "center"),
        (38, "Mean_Sec4", "Mean Sec 4: Tools & Cases (T1-T6)", c_amber, 16, "center"),
        (39, "Mean_Sec5", "Mean Sec 5: Value & Transfer (V1-V6)", c_amber, 16, "center"),
        (40, "Overall_Mean", "Overall Mean Score (1-5)", c_amber, 16, "center"),

        # Key Recommendations (Cols 41-45)
        (41, "REC_1", "6.1: Most relevant findings, trends, tools or cases", c_slate, 40, "left"),
        (42, "REC_2", "6.2: Missing trends, AI types or challenges", c_slate, 40, "left"),
        (43, "REC_3", "6.3: Sections or statements to clarify/update", c_slate, 40, "left"),
        (44, "REC_4", "6.4: Single most important improvement", c_slate, 40, "left"),
        (45, "REC_5", "6.5: Optional additional comments", c_slate, 40, "left"),
    ]

    ws_responses.row_dimensions[2].height = 36

    for col_idx, code, label, color, width, align in col_defs:
        fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        cell = ws_responses.cell(row=2, column=col_idx, value=f"{code}\n({label})" if code != label else code)
        cell.font = font_header
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws_responses.column_dimensions[col_letter].width = width

    # Freeze top 2 rows
    ws_responses.freeze_panes = "A3"

    # Insert a sample illustrative row to test formula syntax and styling
    sample_row = 3
    ws_responses.cell(row=sample_row, column=1, value="2026-09-29 12:00:00").alignment = Alignment(horizontal="center")
    ws_responses.cell(row=sample_row, column=2, value="npc-1").alignment = Alignment(horizontal="center")
    ws_responses.cell(row=sample_row, column=3, value="EN").alignment = Alignment(horizontal="center")
    ws_responses.cell(row=sample_row, column=4, value="Sample NPC Member")
    ws_responses.cell(row=sample_row, column=5, value="npc.expert@industry.eu")
    ws_responses.cell(row=sample_row, column=6, value="Manufacturing Competence Center / SME")
    ws_responses.cell(row=sample_row, column=7, value="Training & Innovation Manager")
    ws_responses.cell(row=sample_row, column=8, value="Spain").alignment = Alignment(horizontal="center")
    ws_responses.cell(row=sample_row, column=9, value="Industrial company").alignment = Alignment(horizontal="left")
    ws_responses.cell(row=sample_row, column=10, value="")

    # Sample Likert scores (Cols 11 to 35: exactly 25 questions)
    sample_scores = [5, 4, 5, 4, 5, 5,   # P1-P6
                     4, 5, 4, 4, 5, 4, 5, # Q1-Q7
                     5, 4, 5, 4, 5, 4,   # T1-T6
                     5, 4, 4, 5, 5, 5]   # V1-V6
    for i, score in enumerate(sample_scores):
        c = ws_responses.cell(row=sample_row, column=11 + i, value=score)
        c.alignment = Alignment(horizontal="center")

    # Formulas for row 3 (standard English Excel formulas)
    # Sec 2: P1-P6 (cols 11-16: K to P)
    ws_responses.cell(row=sample_row, column=36, value=f"=AVERAGE(K{sample_row}:P{sample_row})").number_format = '0.00'
    # Sec 3: Q1-Q7 (cols 17-23: Q to W)
    ws_responses.cell(row=sample_row, column=37, value=f"=AVERAGE(Q{sample_row}:W{sample_row})").number_format = '0.00'
    # Sec 4: T1-T6 (cols 24-29: X to AC)
    ws_responses.cell(row=sample_row, column=38, value=f"=AVERAGE(X{sample_row}:AC{sample_row})").number_format = '0.00'
    # Sec 5: V1-V6 (cols 30-35: AD to AI)
    ws_responses.cell(row=sample_row, column=39, value=f"=AVERAGE(AD{sample_row}:AI{sample_row})").number_format = '0.00'
    # Overall Mean: P1 to V6 (cols 11-35: K to AI)
    ws_responses.cell(row=sample_row, column=40, value=f"=AVERAGE(K{sample_row}:AI{sample_row})").number_format = '0.00'

    for col_f in range(36, 41):
        cell_f = ws_responses.cell(row=sample_row, column=col_f)
        cell_f.alignment = Alignment(horizontal="center")
        cell_f.font = font_formula
        cell_f.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

    # Sample Qualitative Recommendations
    ws_responses.cell(row=sample_row, column=41, value="The concrete industrial case studies in section 4 provide very practical inspiration for SME trainers.").alignment = Alignment(vertical="top", wrap_text=True)
    ws_responses.cell(row=sample_row, column=42, value="Include more focus on open-source edge AI models suitable for manufacturing shop-floors.").alignment = Alignment(vertical="top", wrap_text=True)
    ws_responses.cell(row=sample_row, column=43, value="Clarify the differentiation between generic GenAI tools and specialized training copilots.").alignment = Alignment(vertical="top", wrap_text=True)
    ws_responses.cell(row=sample_row, column=44, value="Provide a one-page executive summary taxonomy of tools at the beginning of the report.").alignment = Alignment(vertical="top", wrap_text=True)
    ws_responses.cell(row=sample_row, column=45, value="Congratulations on the comprehensive mapping; it establishes a solid baseline for the project.").alignment = Alignment(vertical="top", wrap_text=True)

    for c in range(1, 46):
        cell = ws_responses.cell(row=sample_row, column=c)
        cell.border = thin_border
        if cell.font != font_formula:
            cell.font = font_data

    # Auto-filter over all 45 columns
    ws_responses.auto_filter.ref = f"A2:AS{sample_row}"

    # -------------------------------------------------------------
    # 2. SHEET 2: QUESTION CODEBOOK (Codebook_Questions)
    # -------------------------------------------------------------
    ws_codebook = wb.create_sheet(title="Codebook_Questions")
    ws_codebook.views.sheetView[0].showGridLines = True

    # Title Banner
    ws_codebook.merge_cells("A1:E1")
    title_cell = ws_codebook["A1"]
    title_cell.value = "LEARNING BRAINS (ERASMUS+) - NPC 1 CONSULTATION CODEBOOK & VARIABLE DICTIONARY"
    title_cell.font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_codebook.row_dimensions[1].height = 32

    # Headers for Questions Table
    headers_cb = ["Item Code", "Section Name", "Official Question Text (English)", "Question Text (Spanish Reference)", "Scale / Response Type"]
    for c_idx, h_text in enumerate(headers_cb, 1):
        cell = ws_codebook.cell(row=3, column=c_idx, value=h_text)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        cell.alignment = Alignment(horizontal="center" if c_idx in [1, 5] else "left", vertical="center")
        cell.border = thin_border
    ws_codebook.row_dimensions[3].height = 26

    questions_data = [
        # Section 2: R1 purpose and relevance (P1-P6)
        ("P1", "Section 2. R1 purpose and relevance", 
         "The purpose and intended use of the Mapping Report are clear.", 
         "El propósito y el uso previsto del Informe de Mapeo son claros.", 
         "Likert Scale (1 - 5 + N/A)"),
        ("P2", "Section 2. R1 purpose and relevance", 
         "The report addresses relevant needs related to AI-supported on-the-job training in industrial companies.", 
         "El informe responde a necesidades relevantes sobre formación en el puesto de trabajo con IA en empresas industriales.", 
         "Likert Scale (1 - 5 + N/A)"),
        ("P3", "Section 2. R1 purpose and relevance", 
         "The report is relevant to HR managers, training managers, company trainers and other intended target groups.", 
         "El informe es relevante para responsables de RRHH, directores de formación, formadores y otros grupos destinatarios.", 
         "Likert Scale (1 - 5 + N/A)"),
        ("P4", "Section 2. R1 purpose and relevance", 
         "The report provides a useful overview of current developments in AI applied to workplace learning and training.", 
         "El informe ofrece una visión global útil de los avances de la IA aplicada a la formación en el trabajo.", 
         "Likert Scale (1 - 5 + N/A)"),
        ("P5", "Section 2. R1 purpose and relevance", 
         "The report maintains an appropriate focus on industrial companies and SMEs.", 
         "El informe mantiene un enfoque adecuado en empresas industriales y PYMEs.", 
         "Likert Scale (1 - 5 + N/A)"),
        ("P6", "Section 2. R1 purpose and relevance", 
         "The findings are sufficiently relevant to support the development of the Learning Brains Training Programme.", 
         "Las conclusiones son suficientemente relevantes para sustentar el desarrollo del Itinerario Formativo.", 
         "Likert Scale (1 - 5 + N/A)"),

        # Section 3: Quality and coverage of the mapping (Q1-Q7)
        ("Q1", "Section 3. Quality and coverage of the mapping", 
         "The report is clearly structured and easy to navigate.", 
         "El informe está claramente estructurado y es fácil de navegar y consultar.", 
         "Likert Scale (1 - 5 + N/A)"),
        ("Q2", "Section 3. Quality and coverage of the mapping", 
         "The main findings and trends are explained in a clear and understandable way.", 
         "Los hallazgos y tendencias principales se explican de forma clara y comprensible.", 
         "Likert Scale (1 - 5 + N/A)"),
        ("Q3", "Section 3. Quality and coverage of the mapping", 
         "The range of AI tools and applications covered is appropriate for the purpose of the report.", 
         "La gama de herramientas y aplicaciones de IA abordadas es adecuada para el propósito del informe.", 
         "Likert Scale (1 - 5 + N/A)"),
        ("Q4", "Section 3. Quality and coverage of the mapping", 
         "The report provides a balanced view of opportunities, limitations and implementation challenges.", 
         "El informe ofrece una visión equilibrada de oportunidades, limitaciones y desafíos de implantación.", 
         "Likert Scale (1 - 5 + N/A)"),
        ("Q5", "Section 3. Quality and coverage of the mapping", 
         "The information presented appears sufficiently current and relevant to today’s training context.", 
         "La información presentada está suficientemente actualizada y vigente para la formación actual.", 
         "Likert Scale (1 - 5 + N/A)"),
        ("Q6", "Section 3. Quality and coverage of the mapping", 
         "The report reflects a useful range of company, sectoral and/or national contexts.", 
         "El informe refleja una variedad útil de contextos empresariales, sectoriales y nacionales.", 
         "Likert Scale (1 - 5 + N/A)"),
        ("Q7", "Section 3. Quality and coverage of the mapping", 
         "The conclusions are consistent with the evidence and examples presented in the report.", 
         "Las conclusiones son coherentes con las evidencias y ejemplos presentados.", 
         "Likert Scale (1 - 5 + N/A)"),

        # Section 4: AI tools, trends and case studies (T1-T6)
        ("T1", "Section 4. AI tools, trends and case studies", 
         "The report helps readers understand the potential strengths and limitations of different AI applications.", 
         "El informe ayuda a comprender las fortalezas y limitaciones de diferentes aplicaciones de IA.", 
         "Likert Scale (1 - 5 + N/A)"),
        ("T2", "Section 4. AI tools, trends and case studies", 
         "The case studies or practical examples are clear and useful.", 
         "Los casos de estudio o ejemplos prácticos son claros y útiles.", 
         "Likert Scale (1 - 5 + N/A)"),
        ("T3", "Section 4. AI tools, trends and case studies", 
         "The selected examples illustrate realistic applications of AI in workplace learning.", 
         "Los ejemplos seleccionados ilustran aplicaciones realistas de la IA en la formación laboral.", 
         "Likert Scale (1 - 5 + N/A)"),
        ("T4", "Section 4. AI tools, trends and case studies", 
         "The report provides enough information to understand the suitability of different approaches for different organisational contexts.", 
         "El informe aporta información suficiente sobre la idoneidad de los enfoques según el contexto organizativo.", 
         "Likert Scale (1 - 5 + N/A)"),
        ("T5", "Section 4. AI tools, trends and case studies", 
         "The trends and practices highlighted are relevant enough to inform future training content.", 
         "Las tendencias y prácticas destacadas son relevantes para fundamentar futuros contenidos formativos.", 
         "Likert Scale (1 - 5 + N/A)"),
        ("T6", "Section 4. AI tools, trends and case studies", 
         "The report helps distinguish promising practical uses of AI from applications with limited relevance to workplace training.", 
         "El informe ayuda a distinguir usos prometedores de la IA de aplicaciones de escasa utilidad formativa.", 
         "Likert Scale (1 - 5 + N/A)"),

        # Section 5: Practical value and transferability (V1-V6)
        ("V1", "Section 5. Practical value and transferability", 
         "The report can help HR and training professionals identify possible uses of AI in their own organisations.", 
         "El informe ayuda a profesionales de RRHH y formación a identificar posibles usos en sus organizaciones.", 
         "Likert Scale (1 - 5 + N/A)"),
        ("V2", "Section 5. Practical value and transferability", 
         "The findings are transferable to different industrial company contexts.", 
         "Las conclusiones son transferibles a diversos contextos de empresas industriales.", 
         "Likert Scale (1 - 5 + N/A)"),
        ("V3", "Section 5. Practical value and transferability", 
         "The report provides useful reference points for organisations with different levels of digital maturity.", 
         "El informe ofrece puntos de referencia útiles para organizaciones con distintos niveles de madurez digital.", 
         "Likert Scale (1 - 5 + N/A)"),
        ("V4", "Section 5. Practical value and transferability", 
         "The report supports informed decision-making rather than presenting AI tools as solutions in themselves.", 
         "El informe respalda la toma de decisiones informada sin presentar las herramientas como fines en sí mismas.", 
         "Likert Scale (1 - 5 + N/A)"),
        ("V5", "Section 5. Practical value and transferability", 
         "Ethical, legal, data protection and/or inclusion considerations are given appropriate attention where relevant.", 
         "Se presta la debida atención a aspectos éticos, legales, protección de datos e inclusión cuando corresponde.", 
         "Likert Scale (1 - 5 + N/A)"),
        ("V6", "Section 5. Practical value and transferability", 
         "Overall, the report provides a solid knowledge base for the Learning Brains Training Programme and Toolkit.", 
         "En conjunto, el informe aporta una base sólida para el Itinerario Formativo y el Toolkit de Learning Brains.", 
         "Likert Scale (1 - 5 + N/A)"),

        # Section 6: Key recommendations (REC_1 to REC_5)
        ("REC_1", "Section 6. Key recommendations", 
         "Which findings, trends, tools or case studies do you consider most relevant and worth emphasising in the final Mapping Report?", 
         "¿Qué hallazgos, tendencias, herramientas o casos consideras más relevantes y destacables en el informe final?", 
         "Open Qualitative Text"),
        ("REC_2", "Section 6. Key recommendations", 
         "Is there any important trend, type of AI application, implementation challenge or perspective that is missing or underrepresented?", 
         "¿Falta o está poco representada alguna tendencia, tipo de IA, reto de implantación o perspectiva relevante?", 
         "Open Qualitative Text"),
        ("REC_3", "Section 6. Key recommendations", 
         "Are there any sections or statements that should be clarified, corrected, updated or simplified?", 
         "¿Hay secciones o afirmaciones que deban aclararse, corregirse, actualizarse o simplificarse?", 
         "Open Qualitative Text"),
        ("REC_4", "Section 6. Key recommendations", 
         "What is the single most important improvement you would recommend before finalising the Mapping Report?", 
         "¿Cuál es la mejora principal y más importante que recomendarías antes de cerrar el informe?", 
         "Open Qualitative Text"),
        ("REC_5", "Section 6. Key recommendations", 
         "Optional: Is there anything else you would like to add that has not been covered by the previous questions?", 
         "Opcional: ¿Deseas añadir algún otro comentario u observación no abordada anteriormente?", 
         "Open Qualitative Text"),
    ]

    for idx, (code, sec, text_en, text_es, scale) in enumerate(questions_data, 4):
        ws_codebook.cell(row=idx, column=1, value=code).alignment = Alignment(horizontal="center")
        ws_codebook.cell(row=idx, column=2, value=sec)
        ws_codebook.cell(row=idx, column=3, value=text_en).alignment = Alignment(wrap_text=True)
        ws_codebook.cell(row=idx, column=4, value=text_es).alignment = Alignment(wrap_text=True)
        ws_codebook.cell(row=idx, column=5, value=scale).alignment = Alignment(horizontal="center")

        for c in range(1, 6):
            cell = ws_codebook.cell(row=idx, column=c)
            cell.border = thin_border
            cell.font = font_data

    ws_codebook.column_dimensions["A"].width = 16
    ws_codebook.column_dimensions["B"].width = 38
    ws_codebook.column_dimensions["C"].width = 52
    ws_codebook.column_dimensions["D"].width = 52
    ws_codebook.column_dimensions["E"].width = 24

    # Scale definition table in Codebook
    row_scale_start = len(questions_data) + 6
    ws_codebook.cell(row=row_scale_start, column=1, value="Rating Scale Definition:").font = Font(name="Calibri", size=11, bold=True)
    scale_rows = [
        ("1", "Strongly Disagree / Totalmente en desacuerdo"),
        ("2", "Disagree / En desacuerdo"),
        ("3", "Neither agree nor disagree / Neutral (Ni de acuerdo ni en desacuerdo)"),
        ("4", "Agree / De acuerdo"),
        ("5", "Strongly Agree / Totalmente de acuerdo"),
        ("N/A", "Not applicable / I cannot assess / No aplicable o no puedo evaluar"),
    ]
    for i, (val, desc) in enumerate(scale_rows, row_scale_start + 1):
        c1 = ws_codebook.cell(row=i, column=1, value=val)
        c1.alignment = Alignment(horizontal="center")
        c1.font = font_data
        c1.border = thin_border
        c2 = ws_codebook.cell(row=i, column=2, value=desc)
        c2.font = font_data
        c2.border = thin_border

    # -------------------------------------------------------------
    # 3. SHEET 3: EXECUTIVE KPI DASHBOARD (KPI_Dashboard)
    # -------------------------------------------------------------
    ws_kpi = wb.create_sheet(title="KPI_Dashboard")
    ws_kpi.views.sheetView[0].showGridLines = True

    ws_kpi.merge_cells("A1:D1")
    kpi_title = ws_kpi["A1"]
    kpi_title.value = "LEARNING BRAINS (ERASMUS+) - NPC 1 EXECUTIVE EVALUATION DASHBOARD"
    kpi_title.font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    kpi_title.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    kpi_title.alignment = Alignment(horizontal="center", vertical="center")
    ws_kpi.row_dimensions[1].height = 30

    sec_headers = ["Evaluation Dimension / Area", "Key Metric", "Score Range / Target", "Calculation Basis"]
    for c_idx, h in enumerate(sec_headers, 1):
        c = ws_kpi.cell(row=3, column=c_idx, value=h)
        c.font = font_header
        c.fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        c.alignment = Alignment(horizontal="center" if c_idx > 1 else "left")
        c.border = thin_border

    dim_rows = [
        ("Total Completed Validations Received", "=COUNTA(Validation_Responses!A3:A500)", "Total Count", "Respondent entries count"),
        ("Section 2: Purpose & Relevance Mean", "=AVERAGE(Validation_Responses!AJ3:AJ500)", "1.00 - 5.00 (Target >= 4.0)", "Average of items P1 to P6"),
        ("Section 3: Quality & Coverage of Mapping Mean", "=AVERAGE(Validation_Responses!AK3:AK500)", "1.00 - 5.00 (Target >= 4.0)", "Average of items Q1 to Q7"),
        ("Section 4: AI Tools, Trends & Case Studies Mean", "=AVERAGE(Validation_Responses!AL3:AL500)", "1.00 - 5.00 (Target >= 4.0)", "Average of items T1 to T6"),
        ("Section 5: Practical Value & Transferability Mean", "=AVERAGE(Validation_Responses!AM3:AM500)", "1.00 - 5.00 (Target >= 4.0)", "Average of items V1 to V6"),
        ("OVERALL NPC 1 EVALUATION SCORE", "=AVERAGE(Validation_Responses!AN3:AN500)", "1.00 - 5.00 (Target >= 4.0)", "Overall Mean across all 25 Likert items"),
    ]

    for idx, (dim, formula, rng, desc) in enumerate(dim_rows, 4):
        ws_kpi.cell(row=idx, column=1, value=dim).font = font_data
        val_cell = ws_kpi.cell(row=idx, column=2, value=formula)
        val_cell.font = font_formula
        val_cell.alignment = Alignment(horizontal="center")
        val_cell.number_format = '0.00' if idx > 4 else '0'
        ws_kpi.cell(row=idx, column=3, value=rng).alignment = Alignment(horizontal="center")
        ws_kpi.cell(row=idx, column=4, value=desc).alignment = Alignment(horizontal="left")

        if idx == 9:  # Highlight overall mean
            for c in range(1, 5):
                ws_kpi.cell(row=idx, column=c).fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
                ws_kpi.cell(row=idx, column=c).font = Font(name="Calibri", size=10, bold=True, color="92400E")

        for c in range(1, 5):
            ws_kpi.cell(row=idx, column=c).border = thin_border

    # Country Distribution Table
    ws_kpi.cell(row=12, column=1, value="Country Distribution").font = Font(name="Calibri", size=11, bold=True)
    ws_kpi.cell(row=13, column=1, value="Country Name").font = font_header
    ws_kpi.cell(row=13, column=1).fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    ws_kpi.cell(row=13, column=2, value="Responses Count").font = font_header
    ws_kpi.cell(row=13, column=2).fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    ws_kpi.cell(row=13, column=2).alignment = Alignment(horizontal="center")

    countries = ["Spain", "Italy", "Portugal", "Austria", "Slovakia"]
    for i, country in enumerate(countries, 14):
        ws_kpi.cell(row=i, column=1, value=country).font = font_data
        ws_kpi.cell(row=i, column=1).border = thin_border
        count_cell = ws_kpi.cell(row=i, column=2, value=f'=COUNTIF(Validation_Responses!H3:H500, "{country}")')
        count_cell.font = font_data
        count_cell.alignment = Alignment(horizontal="center")
        count_cell.border = thin_border

    # Professional Profile Distribution Table
    ws_kpi.cell(row=21, column=1, value="Professional Profile Distribution").font = Font(name="Calibri", size=11, bold=True)
    ws_kpi.cell(row=22, column=1, value="Profile Category").font = font_header
    ws_kpi.cell(row=22, column=1).fill = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")
    ws_kpi.cell(row=22, column=2, value="Responses Count").font = font_header
    ws_kpi.cell(row=22, column=2).fill = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")
    ws_kpi.cell(row=22, column=2).alignment = Alignment(horizontal="center")

    profiles = [
        ("Industrial company", 'Industrial company'),
        ("HR / training management", 'HR / training management'),
        ("Company trainer / c-VET", 'Company trainer / c-VET'),
        ("AI / digital technologies", 'AI / digital technologies'),
        ("Industry 4.0 / innovation", 'Industry 4.0 / innovation'),
        ("Research / academia", 'Research / academia'),
        ("Other", 'Other*'),
    ]
    for i, (label, search_val) in enumerate(profiles, 23):
        ws_kpi.cell(row=i, column=1, value=label).font = font_data
        ws_kpi.cell(row=i, column=1).border = thin_border
        count_cell = ws_kpi.cell(row=i, column=2, value=f'=COUNTIF(Validation_Responses!I3:I500, "{search_val}")')
        count_cell.font = font_data
        count_cell.alignment = Alignment(horizontal="center")
        count_cell.border = thin_border

    ws_kpi.column_dimensions["A"].width = 46
    ws_kpi.column_dimensions["B"].width = 22
    ws_kpi.column_dimensions["C"].width = 26
    ws_kpi.column_dimensions["D"].width = 38

    # Save
    out_dir = os.path.join(os.path.dirname(__file__), "..", "public", "documents", "validation")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "Learning_Brains_NPC1_Validation_Responses.xlsx")
    wb.save(out_path)
    print(f"Workbook successfully saved to: {out_path}")

if __name__ == "__main__":
    create_npc1_workbook()
